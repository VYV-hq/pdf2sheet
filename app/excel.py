"""Merged Excel workbook + CSV writer (memory-efficient write_only mode)."""
import csv
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

TXN_HEADERS = ["No", "페이지", "계좌번호", "고객명", "상품명", "거래일시", "상태",
               "거래구분", "거래금액", "표면잔액", "취급점", "적요(의뢰인등)",
               "수표번호", "이체상대계좌번호", "검증", "비고"]

COL_WIDTHS = [7, 7, 17, 12, 22, 20, 7, 11, 13, 14, 18, 22, 16, 20, 10, 24]


def _dt(s):
    try:
        return datetime.strptime(s, "%Y/%m/%d %H:%M:%S")
    except (ValueError, TypeError):
        try:
            return datetime.strptime(s, "%Y/%m/%d")
        except (ValueError, TypeError):
            return None


def write_outputs(job_dir: Path, rows, pages, issues, errors):
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9E1F2")
    warn_fill = PatternFill("solid", fgColor="FFC7CE")

    def header_cell(ws, v):
        c = WriteOnlyCell(ws, value=v)
        c.font = bold
        c.fill = fill
        return c

    # ---------------------------------------------------------- transactions
    ws = wb.create_sheet("전체거래내역")
    for i, wdt in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.append([header_cell(ws, h) for h in TXN_HEADERS])

    for i, r in enumerate(rows, 1):
        dt = _dt(r["거래일시"])
        date_cell = WriteOnlyCell(ws, value=dt if dt else r["거래일시"])
        if dt:
            date_cell.number_format = "yyyy-mm-dd hh:mm:ss"
        amt = WriteOnlyCell(ws, value=r["거래금액"])
        bal = WriteOnlyCell(ws, value=r["표면잔액"])
        amt.number_format = bal.number_format = "#,##0"
        chk = WriteOnlyCell(ws, value=r["검증"])
        if not r["검증"].startswith("OK"):
            chk.fill = warn_fill
        ws.append([i, r["페이지"], r["계좌번호"], r["고객명"], r["상품명"],
                   date_cell, r["상태"], r["거래구분"], amt, bal, r["취급점"],
                   r["적요(의뢰인등)"], r["수표번호"], r["이체상대계좌번호"],
                   chk, r["비고"]])

    # --------------------------------------------------------------- summary
    ws2 = wb.create_sheet("계좌별요약")
    for i, wdt in enumerate([17, 12, 22, 12, 10, 16, 16, 16, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = wdt
    ws2.append([header_cell(ws2, h) for h in
                ["계좌번호", "고객명", "상품명", "페이지범위", "거래건수",
                 "입금합계", "지급합계", "최종잔액", "검증률(%)"]])
    stats = {}
    for r in rows:
        s = stats.setdefault(r["계좌번호"], {
            "고객명": r["고객명"], "상품명": r["상품명"],
            "p0": r["페이지"], "p1": r["페이지"], "n": 0,
            "in": 0, "out": 0, "bal": None, "ok": 0})
        s["p0"] = min(s["p0"], r["페이지"]); s["p1"] = max(s["p1"], r["페이지"])
        s["n"] += 1
        if r["거래금액"]:
            if r["상태"] == "입금":
                s["in"] += r["거래금액"]
            elif r["상태"] == "지급":
                s["out"] += r["거래금액"]
        if r["표면잔액"] is not None:
            s["bal"] = r["표면잔액"]
        s["ok"] += r["검증"].startswith("OK")
    for acct, s in stats.items():
        for cell_row in [[acct, s["고객명"], s["상품명"], f"{s['p0']}~{s['p1']}",
                          s["n"], s["in"], s["out"], s["bal"],
                          round(s["ok"] / s["n"] * 100, 1) if s["n"] else 0]]:
            ws2.append(cell_row)

    # ------------------------------------------------------------------- log
    ws3 = wb.create_sheet("처리로그")
    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 100
    ws3.append([header_cell(ws3, "구분"), header_cell(ws3, "내용")])
    kinds = {}
    for p in pages:
        k = p["kind"] + (f"({p['profile']})" if p.get("profile") else "")
        kinds[k] = kinds.get(k, 0) + 1
    ws3.append(["페이지구성", ", ".join(f"{k}: {v}쪽" for k, v in sorted(kinds.items()))])
    for p in pages:
        if p["kind"] == "other" and p.get("text"):
            snippet = re.sub(r"\s+", " ", p["text"])[:300]
            ws3.append([f"p{p['page']}", f"표 없음(문서 페이지): {snippet}"])
        elif p["kind"] == "error":
            ws3.append([f"p{p['page']}", "페이지 처리 오류"])
    for e in errors:
        ws3.append([f"p{e['page']}", f"오류: {e['error'][:400]}"])
    for it in issues:
        ws3.append([f"p{it['페이지']}", f"행{it['행']} {it['계좌번호']}: {it['사유']}"])

    xlsx_path = job_dir / "output.xlsx"
    wb.save(xlsx_path)

    # ------------------------------------------------------------------- csv
    csv_path = job_dir / "output.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(TXN_HEADERS)
        for i, r in enumerate(rows, 1):
            w.writerow([i, r["페이지"], r["계좌번호"], r["고객명"], r["상품명"],
                        r["거래일시"], r["상태"], r["거래구분"], r["거래금액"],
                        r["표면잔액"], r["취급점"], r["적요(의뢰인등)"],
                        r["수표번호"], r["이체상대계좌번호"], r["검증"], r["비고"]])
    return xlsx_path, csv_path
